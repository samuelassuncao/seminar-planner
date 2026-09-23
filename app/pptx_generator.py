from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt

from app.models import FinalSeminarPlan


def generate_pptx(plan: FinalSeminarPlan) -> BytesIO:
    presentation = Presentation()

    # Define explicitamente o formato widescreen 16:9 (13.33" x 7.5")
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    slide_width_in = presentation.slide_width / 914400  # EMU -> polegadas
    margin = 0.7

    for slide_data in plan.slides:
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[5]
        )

        title = slide.shapes.title
        title.text = slide_data.title

        title.left = Inches(0.6)
        title.top = Inches(0.4)
        title.width = presentation.slide_width - Inches(1.2)
        title.height = Inches(0.6)

        for paragraph in title.text_frame.paragraphs:
            paragraph.font.size = Pt(28)

        textbox = slide.shapes.add_textbox(
            Inches(margin),
            Inches(1.2),
            presentation.slide_width - Inches(margin * 2),
            Inches(5.8),
        )

        text_frame = textbox.text_frame
        text_frame.word_wrap = True

        first_paragraph = True

        for section in slide_data.sections:
            if not first_paragraph:
                paragraph = text_frame.add_paragraph()
                paragraph.space_before = Pt(12)

            section_title = (
                text_frame.paragraphs[0]
                if first_paragraph
                else text_frame.add_paragraph()
            )

            section_title.text = section.title
            section_title.font.bold = True
            section_title.font.size = Pt(18)

            first_paragraph = False

            if section.text:
                paragraph = text_frame.add_paragraph()
                paragraph.text = section.text
                paragraph.font.size = Pt(14)
                paragraph.space_before = Pt(4)

            for subtopic in section.subtopics:
                paragraph = text_frame.add_paragraph()
                paragraph.text = f"- {subtopic}"
                paragraph.font.size = Pt(14)
                paragraph.space_before = Pt(3)

    output = BytesIO()
    presentation.save(output)
    output.seek(0)

    return output